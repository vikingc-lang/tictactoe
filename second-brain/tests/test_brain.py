"""End-to-end tests: sync real files, link them, search, learn, create artifacts, serve API/MCP."""

from __future__ import annotations

import asyncio
import json
import threading
import time
import urllib.error
import urllib.request
from pathlib import Path

import docx
import pytest
from pptx import Presentation

from secondbrain.brain import Brain, Enrichment
from secondbrain.config import BrainConfig, SourceConfig
from secondbrain.llm import LLMUnavailable
from secondbrain.render import Deck, Slide


class OfflineLLM:
    """Simulates having no Claude credentials."""

    def text(self, *a, **k):
        raise LLMUnavailable("no credentials")

    def structured(self, *a, **k):
        raise LLMUnavailable("no credentials")


class FakeLLM:
    def __init__(self):
        self.prompts = []

    def text(self, system, prompt, max_tokens=0):
        self.prompts.append(prompt)
        return "# Pricing Strategy Brief\n\n## Summary\n\nAcme prefers **value-based** pricing [1].\n\n- point one\n- point two\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n## Sources\n\n[1] Acme"

    def structured(self, system, prompt, schema, max_tokens=0):
        self.prompts.append(prompt)
        if schema is Deck:
            return Deck(title="Acme Pricing", subtitle="Board update",
                        slides=[Slide(title="Value pricing wins", bullets=["Margin up", "  sub point"],
                                      speaker_notes="See [1]")])
        return Enrichment(summary="About Acme.", tags=["pricing"], entities=["Acme Corp"])


@pytest.fixture
def workspace(tmp_path: Path):
    notes = tmp_path / "notes"
    (notes / "clients").mkdir(parents=True)
    (notes / "Acme Corp.md").write_text(
        "# Acme Corp\n\nAcme is a manufacturing client. Pricing strategy review planned for Q3.\n"
        "Key contact: Jane Rivera. See [[Pricing Playbook]] and [meeting](clients/kickoff.md).\n", encoding="utf-8")
    (notes / "Pricing Playbook.md").write_text(
        "# Pricing Playbook\n\nValue-based pricing beats cost-plus pricing for manufacturing clients. "
        "Pricing strategy should anchor on customer value, willingness to pay and competitive pricing.\n",
        encoding="utf-8")
    (notes / "clients" / "kickoff.md").write_text(
        "# Kickoff meeting\n\nAcme kickoff. Discussed pricing strategy, manufacturing margins and value pricing.\n",
        encoding="utf-8")
    d = docx.Document()
    d.add_heading("Supply Chain Assessment", 1)
    d.add_paragraph("Logistics costs rose 12 percent. Warehouse automation recommended for the manufacturing network.")
    d.save(str(notes / "supply.docx"))
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Market Entry"
    s.placeholders[1].text = "Enter the Brazilian market through a distributor partnership."
    prs.save(str(notes / "market.pptx"))
    (notes / "ignore.bin").write_bytes(b"\x00\x01")

    cfg = BrainConfig(data_dir=tmp_path / "data", output_dir=tmp_path / "out",
                      sources=[SourceConfig("notes", "folder", {"path": str(notes)})])
    cfg.data_dir.mkdir()
    cfg.output_dir.mkdir()
    return cfg, notes


def test_sync_parses_formats_and_is_incremental(workspace):
    cfg, notes = workspace
    brain = Brain(cfg, llm=OfflineLLM())
    stats = {s.source: s for s in brain.sync()}
    assert stats["notes"].added == 5 and not stats["notes"].errors
    kinds = brain.store.stats()["by_kind"]
    assert kinds["note"] == 3 and kinds["document"] == 1 and kinds["presentation"] == 1

    again = {s.source: s for s in brain.sync()}["notes"]
    assert again.added == 0 and again.updated == 0 and again.unchanged == 5


def test_links_explicit_and_related(workspace):
    cfg, _ = workspace
    brain = Brain(cfg, llm=OfflineLLM())
    brain.sync()
    acme = brain.store.find_by_title("Acme Corp")
    related = {n["title"]: n for n in brain.related(acme.id)}
    assert related["Pricing Playbook"]["kind"] == "explicit"     # [[wikilink]]
    assert related["Kickoff meeting"]["kind"] == "explicit"      # markdown relative link
    playbook = brain.store.find_by_title("Pricing Playbook")
    kinds = {n["title"]: n["kind"] for n in brain.related(playbook.id)}
    assert "Kickoff meeting" in kinds                            # similar topic, no explicit link


def test_search_finds_content_inside_office_files(workspace):
    cfg, _ = workspace
    brain = Brain(cfg, llm=OfflineLLM())
    brain.sync()
    assert brain.search("warehouse automation")[0].title == "Supply Chain Assessment"
    assert brain.search("Brazilian distributor")[0].title == "Market Entry"
    assert brain.search("pricing strategy")[0].title in {"Pricing Playbook", "Acme Corp", "Kickoff meeting"}


def test_changes_create_revisions_and_deletions_are_forgotten(workspace):
    cfg, notes = workspace
    brain = Brain(cfg, llm=OfflineLLM())
    brain.sync()
    playbook = notes / "Pricing Playbook.md"
    playbook.write_text("# Pricing Playbook\n\nNew guidance: subscription pricing for services.\n", encoding="utf-8")
    import os, time
    os.utime(playbook, (time.time() + 5, time.time() + 5))
    stats = brain.sync()[0]
    assert stats.updated == 1
    doc = brain.store.find_by_title("Pricing Playbook")
    assert len(brain.store.revisions(doc.id)) == 1
    assert brain.search("subscription")[0].doc_id == doc.id

    (notes / "market.pptx").unlink()
    assert brain.sync()[0].removed == 1
    assert brain.search("Brazilian") == []


def test_unreachable_source_keeps_existing_knowledge(workspace, tmp_path):
    cfg, notes = workspace
    brain = Brain(cfg, llm=OfflineLLM())
    brain.sync()
    cfg.sources[0].options["path"] = str(tmp_path / "unplugged-drive")
    stats = brain.sync("notes")[0]
    assert stats.errors and stats.removed == 0
    assert brain.store.stats()["documents"] == 5


def test_remember_learns_and_links(workspace):
    cfg, _ = workspace
    brain = Brain(cfg, llm=OfflineLLM())
    brain.sync()
    r = brain.remember("Acme agreed to pilot value-based pricing in two plants. See [[Acme Corp]].",
                       title="Acme pricing decision", tags=["decision"])
    assert brain.search("pilot plants")[0].doc_id == r["doc_id"]
    assert "Acme Corp" in {n["title"] for n in brain.related(r["doc_id"])}


def test_ask_offline_returns_cited_passages(workspace):
    cfg, _ = workspace
    brain = Brain(cfg, llm=OfflineLLM())
    brain.sync()
    result = brain.ask("What pricing approach suits manufacturing clients?")
    assert result["mode"] == "extractive" and result["sources"]
    assert "[1]" in result["answer"]


def test_ask_with_claude_passes_numbered_sources(workspace):
    cfg, _ = workspace
    llm = FakeLLM()
    brain = Brain(cfg, llm=llm)
    brain.sync()
    result = brain.ask("pricing strategy?")
    assert result["mode"] == "claude"
    assert "<sources>" in llm.prompts[0] and "[1]" in llm.prompts[0]


def test_create_docx_and_deck_and_learn_from_outputs(workspace):
    cfg, _ = workspace
    brain = Brain(cfg, llm=FakeLLM())
    brain.sync()
    d = brain.create_document("Acme pricing strategy", fmt="docx")
    assert Path(d["path"]).suffix == ".docx" and Path(d["markdown_path"]).exists()
    texts = [p.text for p in docx.Document(d["path"]).paragraphs]
    assert "Pricing Strategy Brief" in texts and "point one" in texts

    deck = brain.create_deck("Acme pricing strategy", slides=3)
    prs = Presentation(deck["path"])
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles[:2] == ["Acme Pricing", "Value pricing wins"] and titles[-1] == "Sources"

    stats = {s.source: s for s in brain.sync()}
    assert stats["_outputs"].added == 3        # md + docx + pptx are now part of the brain


def test_deck_offline_fallback(workspace):
    cfg, _ = workspace
    brain = Brain(cfg, llm=OfflineLLM())
    brain.sync()
    deck = brain.create_deck("pricing strategy", slides=2)
    assert deck["mode"] == "extractive" and Path(deck["path"]).exists()


def test_enrich_links_shared_entities(workspace):
    cfg, _ = workspace
    brain = Brain(cfg, llm=FakeLLM())
    brain.sync()
    assert brain.enrich(limit=50)["enriched"] == 5
    doc = brain.store.find_by_title("Supply Chain Assessment")
    assert doc.summary == "About Acme." and doc.entities == ["Acme Corp"]
    assert any(n["kind"] == "entity" for n in brain.related(doc.id, 50))


def _serve(brain, **kw):
    from secondbrain.api import make_server

    server = make_server(brain, "127.0.0.1", 0, **kw)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server, f"http://127.0.0.1:{server.server_address[1]}"


def _call(base, path, body=None, token=None, raw=False):
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    req = urllib.request.Request(base + path, data=json.dumps(body).encode() if body is not None else None,
                                 headers=headers)
    with urllib.request.urlopen(req) as r:
        data = r.read()
        return data if raw else json.loads(data)


def _wait_for_sync(base, token=None):
    for _ in range(200):
        status = _call(base, "/status", token=token)
        if not status["sync"]["running"]:
            return status
        time.sleep(0.05)
    raise AssertionError("sync did not finish")


def test_rest_api(workspace, monkeypatch):
    cfg, _ = workspace
    monkeypatch.setenv("BRAIN_API_TOKEN", "s3cret")
    brain = Brain(cfg, llm=OfflineLLM())
    server, base = _serve(brain)
    try:
        with pytest.raises(urllib.error.HTTPError):
            _call(base, "/status", token="wrong")
        assert _call(base, "/sync", {}, token="s3cret")["started"] is True
        status = _wait_for_sync(base, "s3cret")
        assert status["stats"]["documents"] == 5 and status["sync"]["last_results"][0]["added"] == 5
        hits = _call(base, "/search?q=warehouse", token="s3cret")
        assert hits[0]["title"] == "Supply Chain Assessment"
        assert _call(base, f"/documents/{hits[0]['doc_id']}", token="s3cret")["kind"] == "document"
        original = _call(base, f"/documents/{hits[0]['doc_id']}/file?token=s3cret", raw=True)
        assert original[:2] == b"PK"                               # the real .docx comes back
        assert _call(base, "/remember", {"text": "Call Jane on Friday"}, token="s3cret")["status"] == "added"
        assert _call(base, "/graph", token="s3cret")["nodes"]
    finally:
        server.shutdown()


def test_web_ui_flow(workspace, tmp_path, monkeypatch):
    cfg, notes = workspace
    monkeypatch.delenv("BRAIN_API_TOKEN", raising=False)
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    cfg.config_path = tmp_path / "brain.toml"
    cfg.config_path.write_text(f'[brain]\ndata_dir = {json.dumps(str(cfg.data_dir))}\n'
                               f'output_dir = {json.dumps(str(cfg.output_dir))}\n', encoding="utf-8")
    cfg.sources = []
    brain = Brain(cfg, llm=OfflineLLM())
    server, base = _serve(brain)
    try:
        html = _call(base, "/", raw=True).decode()
        assert "<title>Second Brain</title>" in html

        # add a folder from the UI: it's written to brain.toml and synced in the background
        with pytest.raises(urllib.error.HTTPError):
            _call(base, "/sources", {"type": "folder", "name": "Bad", "target": str(tmp_path / "missing")})
        assert _call(base, "/sources", {"type": "folder", "name": "My Notes", "target": str(notes)})["syncing"]
        status = _wait_for_sync(base)
        assert status["stats"]["documents"] == 5
        assert "My Notes" in cfg.config_path.read_text()
        assert any(s["name"] == "My Notes" for s in status["configured_sources"])

        # create a deck and download it
        deck = _call(base, "/create/deck", {"topic": "pricing strategy", "slides": 2})
        data = _call(base, deck["download"] + "?download=1", raw=True)
        assert data[:2] == b"PK"
        with pytest.raises(urllib.error.HTTPError):
            _call(base, "/outputs/..%2F..%2Fbrain.toml", raw=True)

        # save a Claude key locally
        assert _call(base, "/settings/claude-key", {"key": "sk-ant-test"})["saved"]
        assert "ANTHROPIC_API_KEY=sk-ant-test" in (cfg.data_dir / "secrets.env").read_text()
        assert _call(base, "/status")["claude_configured"] is True

        # remove the source again
        _call(base, "/sources/remove", {"name": "My Notes"})
        assert "My Notes" not in cfg.config_path.read_text()
    finally:
        server.shutdown()


def test_windows_paths_survive_config_round_trip(tmp_path):
    from secondbrain.config import append_source, load_config

    path = tmp_path / "brain.toml"
    path.write_text(f"[brain]\ndata_dir = {json.dumps(str(tmp_path / 'd'))}\n", encoding="utf-8")
    append_source(path, "OneDrive", "folder", path="C:\\Users\\me\\OneDrive - Contoso\\Clients")
    cfg = load_config(str(path))
    assert cfg.sources[0].options["path"] == "C:\\Users\\me\\OneDrive - Contoso\\Clients"


def test_mcp_server_exposes_tools(workspace):
    from secondbrain.mcp_server import build_server

    cfg, _ = workspace
    brain = Brain(cfg, llm=OfflineLLM())
    brain.sync()
    server = build_server(brain)
    names = {t.name for t in asyncio.run(server.list_tools())}
    assert {"search", "ask", "remember", "related", "create_document", "create_deck", "sync", "status"} <= names


# ---------------------------------------------------------------- email
def _make_email(subject, body, attachments=(), sender="Marco Lindqvist <marco@acme.example>", date="Mon, 21 Sep 2026 10:00:00 +0000"):
    from email.message import EmailMessage

    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"], msg["Date"] = sender, "Jane Rivera <jane@firm.example>", subject, date
    msg["Message-ID"] = f"<{abs(hash(subject))}@acme.example>"
    msg.set_content(body)
    for filename, data, maintype, subtype in attachments:
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=filename)
    return msg.as_bytes()


def _docx_bytes(heading, text):
    import io
    d = docx.Document()
    d.add_heading(heading, 1)
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


DOCX = ("application", "vnd.openxmlformats-officedocument.wordprocessingml.document")


def test_email_files_with_attachments_and_mbox(tmp_path, monkeypatch):
    import mailbox

    mail = tmp_path / "mail"
    mail.mkdir()
    (mail / "budget.eml").write_bytes(_make_email(
        "Pumps pilot budget", "Jane, the pilot budget is approved at EUR 300k. Proposal attached.",
        [("Pilot Proposal.docx", _docx_bytes("Pilot Proposal", "Distributor pricing corridors for pumps."), *DOCX),
         ("logo.png", b"\x89PNG....", "image", "png")]))
    box = mailbox.mbox(str(mail / "archive.mbox"))
    box.add(_make_email("Kickoff agenda", "Agenda: churn analysis and warehouse automation."))
    box.add(_make_email("Minutes", "Tom agreed to join distributor interviews.",
                        [("notes.txt", b"Interview list: Globex, Initech.", "text", "plain")]))
    box.close()

    cfg = BrainConfig(data_dir=tmp_path / "data", output_dir=tmp_path / "out",
                      sources=[SourceConfig("mail", "folder", {"path": str(mail)})])
    cfg.data_dir.mkdir()
    cfg.output_dir.mkdir()
    brain = Brain(cfg, llm=OfflineLLM())
    stats = brain.sync("mail")[0]
    assert stats.errors == [] and stats.added == 5          # 3 emails + docx + txt (the png is skipped)

    email_doc = brain.store.find_by_title("Pumps pilot budget")
    assert email_doc.kind == "email" and "Attachments: Pilot Proposal.docx, logo.png" in email_doc.text
    hit = brain.search("distributor pricing corridors")[0]
    assert hit.title == "Pilot Proposal"                    # text inside the attachment is searchable
    rel = {n["title"]: n["kind"] for n in brain.related(email_doc.id)}
    assert rel["Pilot Proposal"] == "attachment"            # attachment linked to its email
    assert brain.search("Globex Initech")[0].title == "notes"

    # Unchanged mail files are not re-parsed on the next sync.
    import secondbrain.connectors.local as local
    monkeypatch.setattr(local, "parse_file_bytes", lambda *a: (_ for _ in ()).throw(AssertionError("re-parsed")))
    monkeypatch.setattr(local, "iter_mbox", lambda *a: (_ for _ in ()).throw(AssertionError("re-parsed")))
    again = brain.sync("mail")[0]
    assert again.unchanged == 5 and again.removed == 0 and again.errors == []


def test_outlook_msg_without_library_is_a_file_error(tmp_path):
    mail = tmp_path / "mail"
    mail.mkdir()
    (mail / "note.msg").write_bytes(b"\xd0\xcf\x11\xe0 not really")
    (mail / "ok.eml").write_bytes(_make_email("Fine", "All good"))
    cfg = BrainConfig(data_dir=tmp_path / "data", output_dir=tmp_path / "out",
                      sources=[SourceConfig("mail", "folder", {"path": str(mail)})])
    cfg.data_dir.mkdir()
    cfg.output_dir.mkdir()
    stats = Brain(cfg, llm=OfflineLLM()).sync("mail")[0]
    assert stats.added == 1 and len(stats.errors) == 1       # one bad file never stops the rest


class FakeIMAP:
    """Just enough of imaplib.IMAP4_SSL for the connector."""
    mailbox: dict[str, dict[int, bytes]] = {}
    fetches: list[str] = []
    logins: list[tuple] = []

    def __init__(self, host, port, timeout=None):
        assert timeout  # never hang on an unreachable server
        self.host, self.selected = host, None

    def login(self, user, password):
        FakeIMAP.logins.append((user, password))
        return "OK", [b"logged in"]

    def select(self, folder, readonly=False):
        assert readonly
        self.selected = folder.strip('"')
        return ("OK", [b"1"]) if self.selected in self.mailbox else ("NO", [b"no such folder"])

    def response(self, code):
        return code, [b"777"]

    def uid(self, command, *args):
        if command == "SEARCH":
            assert args[1].startswith("SINCE ")
            return "OK", [b" ".join(str(u).encode() for u in sorted(self.mailbox[self.selected]))]
        uid, spec = args
        assert spec == "(BODY.PEEK[])"                      # never marks mail as read
        FakeIMAP.fetches.append(uid)
        return "OK", [(b"1 (UID %s BODY[] {n}" % uid.encode(), self.mailbox[self.selected][int(uid)]), b")"]

    def logout(self):
        return "BYE", []


def test_imap_mailbox_sync(tmp_path, monkeypatch):
    import imaplib

    monkeypatch.setattr(imaplib, "IMAP4_SSL", FakeIMAP)
    monkeypatch.setenv("TEST_MAIL_PW", "app-password")
    FakeIMAP.mailbox = {"INBOX": {
        1: _make_email("SOW signed", "Acme signed the statement of work.",
                       [("SOW.docx", _docx_bytes("Statement of Work", "Phase 1 pricing diagnostic, 8 weeks."), *DOCX)]),
        2: _make_email("Lunch?", "Are you free Thursday?", sender="Tom <tom@acme.example>"),
    }}
    FakeIMAP.fetches, FakeIMAP.logins = [], []
    from secondbrain.config import expand_env
    cfg = BrainConfig(data_dir=tmp_path / "data", output_dir=tmp_path / "out", sources=[SourceConfig(
        "work-mail", "imap", expand_env({"host": "imap.example.com", "username": "jane@firm.example",
                                         "password": "${TEST_MAIL_PW}", "folders": ["INBOX"]}))])
    cfg.data_dir.mkdir()
    cfg.output_dir.mkdir()
    brain = Brain(cfg, llm=OfflineLLM())

    stats = brain.sync("work-mail")[0]
    assert stats.added == 3 and stats.errors == []
    assert FakeIMAP.logins == [("jane@firm.example", "app-password")]
    sow = brain.search("pricing diagnostic")[0]
    assert sow.title == "Statement of Work"
    email_doc = brain.store.find_by_title("SOW signed")
    assert email_doc.kind == "email" and "From: Marco Lindqvist" in email_doc.text
    assert sow.doc_id in {n["doc_id"] for n in brain.related(email_doc.id)}

    # Second sync downloads nothing already known; a new message is picked up; deleted mail is kept.
    FakeIMAP.fetches = []
    del FakeIMAP.mailbox["INBOX"][2]
    FakeIMAP.mailbox["INBOX"][3] = _make_email("Follow-up", "Next steering committee on 10 October.")
    stats = brain.sync("work-mail")[0]
    assert FakeIMAP.fetches == ["3"] and stats.added == 1 and stats.removed == 0
    assert brain.store.find_by_title("Lunch?") is not None


def test_add_mailbox_via_api_keeps_password_secret(tmp_path, monkeypatch):
    import imaplib

    class PickyIMAP(FakeIMAP):
        def login(self, user, password):
            if password != "right-pw":
                raise imaplib.IMAP4.error("AUTHENTICATIONFAILED")
            return super().login(user, password)

    monkeypatch.setattr(imaplib, "IMAP4_SSL", PickyIMAP)
    monkeypatch.delenv("BRAIN_API_TOKEN", raising=False)
    FakeIMAP.mailbox = {"INBOX": {1: _make_email("Hello", "First email")}}
    cfg = BrainConfig(data_dir=tmp_path / "data", output_dir=tmp_path / "out")
    cfg.data_dir.mkdir()
    cfg.output_dir.mkdir()
    cfg.config_path = tmp_path / "brain.toml"
    cfg.config_path.write_text(f"[brain]\ndata_dir = {json.dumps(str(cfg.data_dir))}\n"
                               f"output_dir = {json.dumps(str(cfg.output_dir))}\n", encoding="utf-8")
    server, base = _serve(Brain(cfg, llm=OfflineLLM()))
    try:
        req = {"type": "imap", "name": "Personal mail", "target": "me@gmail.com", "password": "wrong"}
        with pytest.raises(urllib.error.HTTPError) as err:
            _call(base, "/sources", req)
        assert "couldn't sign in to imap.gmail.com" in json.loads(err.value.read())["error"]
        assert "Personal mail" not in cfg.config_path.read_text()

        assert _call(base, "/sources", {**req, "password": "right-pw"})["syncing"]
        status = _wait_for_sync(base)
        assert status["stats"]["documents"] == 1
        toml = cfg.config_path.read_text()
        assert "right-pw" not in toml and "${MAIL_PASSWORD_PERSONAL_MAIL}" in toml
        assert "MAIL_PASSWORD_PERSONAL_MAIL=right-pw" in (cfg.data_dir / "secrets.env").read_text()
        mail_src = next(s for s in status["configured_sources"] if s["name"] == "Personal mail")
        assert "password" not in mail_src and mail_src["host"] == "imap.gmail.com"
    finally:
        server.shutdown()
        import os
        os.environ.pop("MAIL_PASSWORD_PERSONAL_MAIL", None)


# ---------------------------------------------------------------- create from chosen references
class ExplodingLLM:
    def text(self, *a, **k):
        raise AssertionError("Claude must not be called")

    structured = text


def test_create_with_chosen_references_uses_only_those(workspace):
    cfg, _ = workspace
    llm = FakeLLM()
    brain = Brain(cfg, llm=llm)
    brain.sync()
    supply = brain.store.find_by_title("Supply Chain Assessment")
    market = brain.store.find_by_title("Market Entry")
    result = brain.create_document("Operations brief", fmt="docx", doc_ids=[supply.id, market.id, supply.id])
    assert result["mode"] == "claude"
    assert [s["title"] for s in result["sources"]] == ["Supply Chain Assessment", "Market Entry"]
    prompt = llm.prompts[-1]
    assert "Warehouse automation" in prompt and "Brazilian" in prompt
    assert "Pricing Playbook" not in prompt and "Acme Corp" not in prompt   # nothing else sneaks in

    deck = brain.create_deck("Operations", slides=2, doc_ids=[market.id])
    assert [s["title"] for s in deck["sources"]] == ["Market Entry"]


def test_create_without_ai_never_calls_claude(workspace):
    cfg, _ = workspace
    brain = Brain(cfg, llm=ExplodingLLM())
    brain.sync()
    playbook = brain.store.find_by_title("Pricing Playbook")
    doc = brain.create_document("Pricing note", fmt="docx", doc_ids=[playbook.id], use_ai=False)
    assert doc["mode"] == "no_ai"
    texts = "\n".join(p.text for p in docx.Document(doc["path"]).paragraphs)
    assert "Value-based pricing beats cost-plus" in texts and "no AI" in texts
    deck = brain.create_deck("Pricing", slides=3, use_ai=False)          # no picks: brain chooses sources
    assert deck["mode"] == "no_ai" and deck["sources"]
    slides = Presentation(deck["path"]).slides
    assert any("Value-based pricing" in sh.text_frame.text for s in slides for sh in s.shapes if sh.has_text_frame)


def test_create_api_accepts_references_and_no_ai(workspace, monkeypatch):
    cfg, _ = workspace
    monkeypatch.delenv("BRAIN_API_TOKEN", raising=False)
    brain = Brain(cfg, llm=ExplodingLLM())
    brain.sync()
    market = brain.store.find_by_title("Market Entry")
    server, base = _serve(brain)
    try:
        r = _call(base, "/create/deck", {"topic": "Brazil", "slides": 2, "doc_ids": [market.id], "use_ai": False})
        assert r["mode"] == "no_ai" and [s["title"] for s in r["sources"]] == ["Market Entry"]
        assert _call(base, r["download"] + "?download=1", raw=True)[:2] == b"PK"
        with pytest.raises(urllib.error.HTTPError):                       # unknown ids -> clear error
            _call(base, "/create/document", {"topic": "x", "doc_ids": [99999], "use_ai": False})
    finally:
        server.shutdown()
