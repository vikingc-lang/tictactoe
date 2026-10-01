"""Turn raw bytes of many file formats into plain text.

Supported out of the box: text/markdown/code, HTML, CSV/JSON, email (.eml; Outlook .msg with extract-msg),
PDF, Word (.docx), PowerPoint (.pptx) and, if ``openpyxl`` is installed, Excel (.xlsx).
"""

from __future__ import annotations

import io
import json
from html.parser import HTMLParser

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".rst", ".org", ".log", ".csv", ".tsv", ".yaml", ".yml", ".toml",
    ".ini", ".py", ".js", ".ts", ".java", ".go", ".rb", ".sql", ".sh", ".tex",
}
RICH_EXTS = {".html", ".htm", ".json", ".eml", ".msg", ".mbox", ".pdf", ".docx", ".pptx", ".xlsx"}
EMAIL_EXTS = {".eml", ".msg", ".mbox"}
SUPPORTED_EXTS = TEXT_EXTS | RICH_EXTS

# One document never contributes more than this much text (about a 600-page book). Giant data exports would
# otherwise take minutes to split and index, and drown out real documents in search.
MAX_TEXT_CHARS = 1_500_000
MAX_SHEET_ROWS = 20_000

KIND_BY_EXT = {
    ".md": "note", ".markdown": "note", ".txt": "note", ".org": "note",
    ".pdf": "pdf", ".docx": "document", ".pptx": "presentation", ".xlsx": "spreadsheet",
    ".csv": "data", ".tsv": "data", ".json": "data", ".html": "web", ".htm": "web", ".eml": "email", ".msg": "email",
}


def kind_for(ext: str) -> str:
    return KIND_BY_EXT.get(ext.lower(), "code" if ext.lower() in TEXT_EXTS else "file")


def _decode(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.title = ""
        self._skip = 0
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag == "title":
            self._in_title = True
        elif tag in self.BLOCK:
            self.parts.append("\n")
        if tag in {"h1", "h2", "h3"}:
            self.parts.append("#" * int(tag[1]) + " ")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        elif tag == "title":
            self._in_title = False
        elif tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self._skip:
            self.parts.append(data)


def html_to_text(html: str) -> tuple[str, str]:
    p = _HTMLText()
    p.feed(html)
    lines = [" ".join(line.split()) for line in "".join(p.parts).splitlines()]
    return p.title.strip(), "\n".join(line for line in lines if line)


def parse(data: bytes, ext: str) -> tuple[str, str | None]:
    """Return (text, title_hint). Raises ValueError for unsupported formats."""
    ext = ext.lower()
    if ext in TEXT_EXTS:
        return _decode(data), None
    if ext in {".html", ".htm"}:
        title, text = html_to_text(_decode(data))
        return text, title or None
    if ext == ".json":
        try:
            return json.dumps(json.loads(_decode(data)), indent=1, ensure_ascii=False), None
        except json.JSONDecodeError:
            return _decode(data), None
    if ext in {".eml", ".msg"}:
        from .emails import parse_file_bytes

        parsed = parse_file_bytes(data, ext)
        return parsed.as_text(), parsed.subject or None
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [f"[page {i + 1}]\n{(page.extract_text() or '').strip()}" for i, page in enumerate(reader.pages)]
        title = (reader.metadata.title if reader.metadata else None) or None
        return "\n\n".join(pages), title
    if ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = []
        for para in d.paragraphs:
            if not para.text.strip():
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            if style.startswith("heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                parts.append("#" * int(level) + " " + para.text)
            else:
                parts.append(para.text)
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts), (d.core_properties.title or None)
    if ext == ".pptx":
        from pptx import Presentation

        prs = Presentation(io.BytesIO(data))
        parts = []
        first_title = None
        for i, slide in enumerate(prs.slides, start=1):
            if first_title is None and slide.shapes.title is not None and slide.shapes.title.text.strip():
                first_title = slide.shapes.title.text.strip()
            texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text]
            parts.append(f"## Slide {i}\n" + "\n".join(texts))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                parts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        core_title = prs.core_properties.title
        if core_title in {"", "PowerPoint Presentation"}:
            core_title = None
        return "\n\n".join(parts), core_title or first_title
    if ext == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:  # optional dependency
            raise ValueError("install openpyxl to index .xlsx files") from exc
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        chars = 0
        for ws in wb.worksheets:
            parts.append(f"## Sheet {ws.title}")
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= MAX_SHEET_ROWS or chars >= MAX_TEXT_CHARS:  # huge data exports: index the top, skip the rest
                    parts.append(f"… (sheet truncated after {n} rows)")
                    break
                if any(v is not None for v in row):
                    line = " | ".join("" if v is None else str(v) for v in row)
                    chars += len(line)
                    parts.append(line)
        return "\n".join(parts), None
    raise ValueError(f"unsupported file type: {ext}")
